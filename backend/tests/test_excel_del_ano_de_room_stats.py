# -*- coding: utf-8 -*-
"""El año entero de estadística del PMS, en un solo Excel.

Owner, 2026-09-11: *«como se podria construir un excel donde se pueda bajar
todo esta informacion por mes, un tab por mes, con todos los sub tabs y uno
final donde este consolidado y se vaya actualizando conforme se van agregando
mas meses»*.

El armado vive en `frontend/lib/roomStatsExcel.ts` porque la pantalla ya tiene
el año entero en memoria (`/anio/`) y así las hojas salen de los MISMOS
números que se están viendo. Acá se blinda lo que puede romperse en silencio:
la aritmética de columnas que el endpoint exige, y que el escritor aguante la
forma que el frontend manda.
"""
import pathlib
import re

import pytest
from openpyxl import load_workbook

from app.export.cuadro_excel import build_cuadros_workbook

FRONT = pathlib.Path(__file__).resolve().parents[2] / "frontend"
FUENTE = FRONT / "lib/roomStatsExcel.ts"
PAGINA = FRONT / "app/month-end/room-stats/page.tsx"


def test_existe_y_la_pantalla_lo_usa():
    assert FUENTE.exists()
    pag = PAGINA.read_text(encoding="utf-8")
    assert 'from "@/lib/roomStatsExcel"' in pag
    assert "async function bajarElAnio" in pag
    assert "⬇ Excel del año" in pag
    # Y el del mes sigue existiendo: son dos necesidades distintas.
    assert "⬇ Excel del mes" in pag


def test_cada_fila_tiene_una_celda_menos_que_columnas():
    """⚠️ La regla de oro del exportador: la PRIMERA columna es la etiqueta,
    así que a `valores` le tocan `len(columnas) - 1` celdas.

    Una fila con una de más el endpoint la rechaza con 422; una de menos se
    corre sola y las cifras quedan bajo el mes equivocado sin que nada avise.
    Acá se verifica la aritmética de las dos hojas leyendo la fuente.
    """
    src = FUENTE.read_text(encoding="utf-8")

    # Hoja del mes: columnas = Canal + categorías + TOTAL.
    assert '{ label: "Canal", ancho: 34, formato: "texto" },' in src
    assert '{ label: "TOTAL", ancho: 18, formato: "num" },' in src
    # …y los valores son categorías + el total. Nunca una categoría suelta.
    assert "valores: [...v, v.reduce((a, x) => a + x, 0)]" in src
    assert "const vacias = cats.map(() => null).concat([null]);" in src

    # Hoja acumulada: columnas = Concepto + 12 meses + YTD.
    assert '{ label: "Concepto", ancho: 38, formato: "texto" },' in src
    assert "...anio.meses.map(m => ({ label: MES3[m.month - 1]" in src
    assert "const vacias = anio.meses.map(() => null).concat([null]);" in src
    # Todo lo que se empuja usa `porMes(...)` + el YTD: doce y uno.
    for patron in (r"\[\.\.\.porMes\([^\]]+\), ytd", r"\[\.\.\.porMes\(dispMes\), ytdDisp\]"):
        assert re.search(patron, src), f"no se encontró {patron}"


def test_un_mes_sin_cargar_va_en_blanco_y_no_en_cero():
    """⚠️ El error más caro de un acumulado de doce columnas.

    Un cero dice «ese mes el hotel no vendió». Un blanco dice «todavía no
    subimos el archivo». Con seis de doce meses cargados, confundirlos mete
    seis meses de ceros en el promedio del año.

    Es la misma regla que `anio_room_stats` ya defiende con `cargado: false`.
    """
    src = FUENTE.read_text(encoding="utf-8")
    assert "(m.cargado ? f(m) : null)" in src, "los meses sin cargar no salen en blanco"
    # Y el YTD suma sólo los cargados.
    assert "const cargados = anio.meses.filter(m => m.cargado);" in src
    assert "cargados.reduce<T3>((a, m) => mas(a, f(m)), cero())" in src


def test_el_ytd_de_una_tasa_se_recalcula_y_no_se_promedia():
    """⚠️ El ADR del año es ingreso acumulado / noches acumuladas.

    Promediar los doce ADR mensuales hace pesar igual a un mes de 20 noches y
    a uno de 202. En Amarena 2026 son $290 contra $255 — un 14% de diferencia
    en el indicador que más se mira, sin ningún síntoma.
    """
    src = FUENTE.read_text(encoding="utf-8")
    assert "adr(ytd(baseMes))" in src
    assert "adr(ytd(m => baseCat(m, c)))" in src
    # Nada que promedie celdas.
    assert "/ 12" not in src
    assert "cargados.length" not in src.split("const ytd =")[1].split("\n\n")[0]


def test_la_hoja_del_mes_trae_las_tres_vistas():
    """Una hoja es UNA tabla, así que las tres vistas de la pantalla no entran
    como tres cuadros. No hace falta: la matriz canal × categoría las contiene
    —sumar una fila da «Por canal», sumar una columna da «Por habitación»— y
    se abre en bloques para que cada uno tenga una sola unidad."""
    src = FUENTE.read_text(encoding="utf-8")
    for bloque in ('bloque("NOCHES OCUPADAS", 0, "num")',
                   'bloque("PAX", 1, "num")',
                   'bloque("INGRESO", 2, "usd2")',
                   'banda("ADR")',
                   'banda("INVENTARIO E INDICADORES")'):
        assert bloque in src, f"falta el bloque {bloque}"
    # El inventario no se abre por canal: es del hotel, no de quien vendió.
    assert '"Noches disponibles"' in src and '"% Ocupación"' in src and '"RevPAR"' in src


def test_el_excel_del_ano_usa_la_misma_base_que_la_pantalla():
    """Un Excel que no coincide con lo que el otro está viendo es peor que no
    tenerlo: se manda por correo y discute contra la pantalla."""
    pag = PAGINA.read_text(encoding="utf-8")
    cuerpo = pag[pag.index("async function bajarElAnio"):pag.index("// ── el calce")]
    assert "enAdr[c.canal_code] ?? c.cuenta_para_kpis" in cuerpo, \
        "el Excel del año no respeta los canales destildados"
    src = FUENTE.read_text(encoding="utf-8")
    assert "Con todos los canales (PDF)" in src, "no se puede cuadrar contra el PMS"


def test_el_canal_fuera_de_la_base_se_rotula_en_el_archivo():
    """En la pantalla el canal excluido sale en itálica con una marca. Un
    Excel no lleva esa marca: si no lo dice el texto, la hoja suelta no
    explica por qué el TOTAL no es la suma de las filas de arriba."""
    src = FUENTE.read_text(encoding="utf-8")
    assert "fuera de la base" in src


# ─────────── Que el escritor aguante la forma que manda el frontend ────────

def _cuadro_mes():
    """La forma exacta de una hoja de mes: bandas sin valores, bloques con
    `formato` por fila, y una fila de porcentaje como fracción."""
    cats = ["Garden View", "Beachfront", "Master Suite"]
    vacias = [None, None, None, None]
    return {
        "titulo": "Agosto 2026 · Estadística de habitaciones",
        "subtitulo": "ACTUAL Final 2026 · canales en filas",
        "hoja": "Ago 2026",
        "columnas": [{"label": "Canal", "ancho": 34, "formato": "texto"}]
                    + [{"label": c, "ancho": 20, "formato": "num"} for c in cats]
                    + [{"label": "TOTAL", "ancho": 18, "formato": "num"}],
        "filas": [
            {"label": "NOCHES OCUPADAS", "es_total": True, "valores": vacias},
            {"label": "DIRECTOS", "nivel": 1, "formato": "num", "valores": [10, 5, 2, 17]},
            {"label": "CPL · fuera de la base", "nivel": 1, "formato": "num",
             "valores": [3, 0, 0, 3]},
            {"label": "TOTAL (base de indicadores)", "es_total": True,
             "formato": "num", "valores": [10, 5, 2, 17]},
            {"label": "Con todos los canales (PDF)", "nivel": 1, "formato": "num",
             "valores": [13, 5, 2, 20]},
            {"label": "INVENTARIO E INDICADORES", "es_total": True, "valores": vacias},
            {"label": "% Ocupación", "nivel": 1, "formato": "pct",
             "valores": [0.4073, 0.31, 0.12, 0.2903]},
            {"label": "RevPAR", "nivel": 1, "formato": "usd2",
             "valores": [103.93, 88.1, 40.0, 77.65]},
        ],
    }


def _cuadro_acumulado():
    """La hoja consolidada: doce meses + YTD, con los meses sin cargar en
    blanco — que es lo que hace distinta a esta hoja de cualquier otra."""
    meses = ["Ene", "Feb", "Mar", "Abr", "May", "Jun",
             "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
    vacias = [None] * 13
    return {
        "titulo": "Acumulado 2026 · Estadística de habitaciones",
        "hoja": "Acumulado 2026",
        "columnas": [{"label": "Concepto", "ancho": 38, "formato": "texto"}]
                    + [{"label": m, "ancho": 14, "formato": "num"} for m in meses]
                    + [{"label": "YTD Ago", "ancho": 16, "formato": "num"}],
        "filas": [
            {"label": "NOCHES OCUPADAS · POR CATEGORÍA", "es_total": True,
             "valores": vacias},
            {"label": "Garden View", "nivel": 1, "formato": "num",
             "valores": [None, None, 20, 46, 60, 63, 132, 202, None, None, None, None, 523]},
            {"label": "INDICADORES DEL HOTEL", "es_total": True, "valores": vacias},
            {"label": "ADR", "nivel": 1, "formato": "usd2",
             "valores": [None, None, 317.66, 355.19, 328.66, 310.70, 273.54, 255.20,
                         None, None, None, None, 281.15]},
        ],
    }


def test_el_escritor_arma_una_hoja_por_mes_mas_el_consolidado():
    """Seis meses + el acumulado son siete pestañas, con los nombres que el
    frontend pide y sin que ninguno se pise."""
    cuadros = [dict(_cuadro_mes(), hoja=f"{m} 2026") for m in
               ("Mar", "Abr", "May", "Jun", "Jul", "Ago")] + [_cuadro_acumulado()]
    wb = load_workbook(filename=__import__("io").BytesIO(
        build_cuadros_workbook(cuadros)))
    assert wb.sheetnames[-1] == "Acumulado 2026"
    assert len([h for h in wb.sheetnames if h.endswith("2026")]) >= 7
    assert "Ago 2026" in wb.sheetnames


def test_los_valores_van_como_numero_y_el_mes_vacio_queda_vacio():
    """⚠️ Nunca `"$1,234.00"` como cadena: un Excel con las cifras en texto no
    se puede sumar ni graficar, que es para lo que se baja. Y la celda de un
    mes sin cargar tiene que quedar REALMENTE vacía, no en cero."""
    import io

    wb = load_workbook(filename=io.BytesIO(
        build_cuadros_workbook([_cuadro_acumulado()])))
    ws = wb["Acumulado 2026"]
    fila = next(r for r in ws.iter_rows()
                if r[0].value == "Garden View")
    # Enero y febrero no están cargados.
    assert fila[1].value is None and fila[2].value is None
    # Marzo sí, y es un número.
    assert fila[3].value == 20
    assert isinstance(fila[3].value, (int, float))
    # El YTD es la suma de lo cargado, no el promedio.
    assert fila[13].value == 523


def test_el_porcentaje_viaja_como_fraccion():
    """⚠️ El formato de Excel es `0.0%`, que multiplica por 100 al mostrar.
    Mandar 40.73 en vez de 0.4073 imprime «4073.0%» — se ve roto, que es lo
    bueno; el riesgo real es mandarlo ya formateado como texto y que nadie
    pueda sumarlo."""
    import io

    wb = load_workbook(filename=io.BytesIO(build_cuadros_workbook([_cuadro_mes()])))
    ws = wb["Ago 2026"]
    fila = next(r for r in ws.iter_rows() if r[0].value == "% Ocupación")
    assert fila[1].value == pytest.approx(0.4073)
    assert "%" in fila[1].number_format


def test_una_banda_sin_valores_no_rompe_el_escritor():
    """Las bandas de sección van con `es_total` y todas las celdas en `None`:
    es lo que separa los bloques de una hoja que trae cuatro unidades."""
    import io

    wb = load_workbook(filename=io.BytesIO(build_cuadros_workbook([_cuadro_mes()])))
    ws = wb["Ago 2026"]
    banda = next(r for r in ws.iter_rows() if r[0].value == "NOCHES OCUPADAS")
    assert banda[0].font.bold
    assert all(c.value is None for c in banda[1:4])
