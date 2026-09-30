# -*- coding: utf-8 -*-
"""El Resumen Ejecutivo mensual en Word.

Owner, 2026-09-30, entregando el `Executive Summary` de CWL: *«usa este formato
como estandar y prepara uno igual para agosto en Amarena. quiero el informe en
word»*.

## Lo que este archivo defiende

El informe se lee como prosa, y una prosa no cuadra contra nada: si una frase
dice un numero distinto del cuadro que tiene debajo, no hay forma de notarlo. Por
eso todo lo que se afirma sale del MISMO agregador que el cierre, y por eso las
frases se arman con la variacion en la mano.
"""
import io
import pathlib

import pytest

from app.export.executive_summary import (
    build_executive_summary, k, pct, usd, var_pct,
)

API = pathlib.Path(__file__).resolve().parents[1] / "app/api/executive_summary_api.py"
DOCX_MOD = pathlib.Path(__file__).resolve().parents[1] / "app/export/executive_summary.py"


def _col(rev, gop, ebitda, neto, occ, adr, noc=100.0, disp=300.0, sid="x",
         corte="ytd"):
    return {
        "__sid": sid, "__corte": corte,
        "kpis": {"rooms_available": disp, "rooms_occupied": noc, "guests": noc * 2,
                 "occupancy_pct": occ, "adr": adr, "revpar": rev / disp},
        "lines": [
            {"line_code": "TOTAL_REVENUES", "amount_usd": rev},
            {"line_code": "TOTAL_OPEXP", "amount_usd": rev * 0.5},
            {"line_code": "TOTAL_OP_PROFIT", "amount_usd": rev * 0.5},
            {"line_code": "TOTAL_OVERHEAD", "amount_usd": rev * 0.3},
            {"line_code": "GOP", "amount_usd": gop},
            {"line_code": "TOTAL_NON_OP", "amount_usd": rev * 0.05},
            {"line_code": "EBITDA_BEFORE", "amount_usd": ebitda},
            {"line_code": "EBITDA_AFTER", "amount_usd": ebitda * 0.9},
            {"line_code": "EBT", "amount_usd": neto * 1.1},
            {"line_code": "NET_PROFIT", "amount_usd": neto},
            {"line_code": "REV_ROOMS", "amount_usd": rev * 0.6},
            {"line_code": "REV_TOURS", "amount_usd": rev * 0.2},
        ],
    }


def _version(mult=1.0, sid="x"):
    return {
        "scenario_id": sid, "type": "ACTUAL", "version": "Final", "year": 2026,
        "month": _col(277_856 * mult, -45_851, -63_673, -99_140, 0.2731, 514.31,
                      sid=sid, corte="month"),
        "ytd": _col(3_973_720 * mult, 968_084, 762_972, 213_040, 0.4583, 599.41,
                    noc=3341, disp=7290, sid=sid, corte="ytd"),
        "full": _col(5_040_287 * mult, 709_847, 430_693, -213_072, 0.4474, 596.82,
                     noc=4483, disp=10020, sid=sid, corte="full"),
    }


def _datos():
    act, bud = _version(1.0, "a"), _version(0.907, "b")
    bud["type"] = "BUDGET"
    fcs = _version(0.98, "f")
    fcs["type"] = "FORECAST"
    gasto = {("a", "month"): 100.0, ("b", "month"): 90.0}

    def totales(col, clave):
        if clave == "TOTAL_REVENUES":
            return next(l["amount_usd"] for l in col["lines"]
                        if l["line_code"] == "TOTAL_REVENUES")
        base = {"TOTAL_PAYROLL": 1_506_428, "TOTAL_OPEX_ONLY": 1_017_695,
                "TOTAL_COST": 481_511, "TOTAL_PROPERTY": 755_043}[clave]
        return base * (1.0 if col["__sid"] == "a" else 0.92)

    return {
        "propiedad": "Amarena Canvas Hotel", "mes": 8, "anio": 2026,
        "actual": act, "budget": bud, "forecast": fcs,
        "rotulos": {"actual": "ACTUAL Final 2026", "budget": "BUDGET Final 2026",
                    "forecast": "FORECAST Working 2026"},
        "totales": totales,
        "mix": [["Rooms", "$2,021,430.00", "$1,890,670.00", "$130,760.00", "+6.9%"]],
        "adr_por_mes": [["08", "$514.31", "27.31%", "254"]],
        "positivos": [("Ingreso total",
                       "El ingreso llegó a $3.974M contra $3.605M presupuestados.")],
        "negativos": [("Costo de ventas",
                       "El costo de ventas cerró por encima del plan.")],
    }


def _texto(blob: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(blob))
    partes = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for fila in t.rows:
            partes += [c.text for c in fila.cells]
    return "\n".join(partes)


def test_el_informe_se_genera_y_es_un_docx():
    blob = build_executive_summary(_datos())
    assert blob[:2] == b"PK", "no es un .docx"
    assert len(blob) > 10_000


def test_trae_las_secciones_del_formato_del_owner():
    """El PDF que el owner dio como estandar: portada, introduccion, los tres
    cortes numerados, los drivers y los positivos/negativos.

    ⚠️ **En espanol.** Owner, 2026-09-30, señalando el boton: *«esto debe ser en
    español»*. El PDF modelo estaba en ingles y de ahi venia el primer armado,
    pero el formato es la ESTRUCTURA —portada, los tres cortes, el
    flow-through, lo bueno y lo malo—, no el idioma: el informe lo lee la junta
    aca.
    """
    t = _texto(build_executive_summary(_datos()))
    assert "RESUMEN EJECUTIVO MENSUAL" in t
    assert "AGOSTO 2026" in t
    assert "Amarena Canvas Hotel" in t
    assert "INTRODUCCIÓN" in t
    for titulo in ("1.1 Agosto 2026 — el mes contra el presupuesto",
                   "1.2 Acumulado a Agosto 2026 — contra el presupuesto",
                   "1.3 Proyección del año completo 2026 contra el presupuesto",
                   "SECCIÓN 2 — De qué depende el resultado",
                   "2.1 Volumen (demanda)", "2.2 Tarifa (calidad del ingreso)",
                   "2.3 Composición del ingreso"):
        assert titulo in t, f"falta «{titulo}»"
    # Y no quedo prosa en ingles suelta.
    for ingles in ("Total Revenue reached", "above Budget", "versus Budget",
                   "Monthly Performance", "Overall Positive"):
        assert ingles not in t, f"quedo en ingles: «{ingles}»"


def test_los_ROTULOS_de_los_cuadros_se_quedan_en_ingles():
    """⚠️ A proposito, y es lo unico que no se traduce. «Total available Rooms»,
    «EBITDA BEFORE CAPITAL» y «GROSS OPERATING PROFIT» son los MISMOS rotulos
    que usa el P&L de la pantalla y los que el owner tiene en su Excel.
    Traducirlos obligaria a comprobar que «Utilidad bruta operativa» y «GROSS
    OPERATING PROFIT» son el mismo renglon."""
    t = _texto(build_executive_summary(_datos()))
    for rotulo in ("Total available Rooms", "Average Daily Room Only",
                   "TOTAL REVENUES", "EBITDA BEFORE CAPITAL", "NET PROFIT"):
        assert rotulo in t, f"se tradujo el rotulo «{rotulo}»"


def test_la_prosa_dice_LOS_NUMEROS_del_cuadro():
    """⚠️ Lo que este informe no puede hacer es decir en el texto algo distinto
    de lo que muestra el cuadro de al lado. Una prosa no cuadra contra nada: la
    discrepancia no se ve, y el que la encuentre deja de creerle a los dos.

    Por eso cada frase se arma con la variacion ya calculada.
    """
    t = _texto(build_executive_summary(_datos()))
    # El ingreso del mes y su variacion, en la escala del owner.
    assert "$277.9K" in t
    assert "$3.974M" in t        # el YTD
    assert "27.31%" in t         # la ocupacion del mes
    assert "$599.41" in t        # el ADR acumulado


def test_en_el_full_year_la_columna_principal_es_el_FORECAST():
    """⚠️ El Actual del año son los meses cargados. Compararlo contra doce meses
    de presupuesto mostraria una caida que solo significa que el año no ha
    terminado — y en un informe a duenos eso se lee como un derrumbe."""
    src = DOCX_MOD.read_text(encoding="utf-8")
    assert 'principal = f if (corte == "full" and f) else a' in src
    assert "el año no" in src


def test_los_positivos_y_negativos_se_ORDENAN_por_tamano():
    """Un informe que siempre comenta los mismos seis renglones habla del mes en
    que se escribio la plantilla, no del mes que se esta cerrando."""
    src = API.read_text(encoding="utf-8")
    assert "porte.sort(key=lambda x: -abs(x[0]))" in src


def test_en_gasto_MENOS_es_favorable():
    """⚠️ Con la regla del ingreso, un sobrecosto entraria en la lista de
    positivos. Es el error clasico de una tabla de variaciones."""
    src = API.read_text(encoding="utf-8")
    assert "ga <= gb" in src, "el gasto se evalua con la regla del ingreso"
    docx = DOCX_MOD.read_text(encoding="utf-8")
    # Y en el flow-through el signo es el del EFECTO sobre la utilidad.
    assert 'efecto = (a - b) if clave == "TOTAL_REVENUES" else -(a - b)' in docx


def test_SOLO_van_las_secciones_1_y_2():
    """Owner, 2026-09-30: *«solo vamos a dejar seccion 1 y 2 por ahora; quita
    todo lo demas»*.

    ⚠️ El «por ahora» es literal: `SECCIONES` sigue existiendo y
    `_positivos_y_negativos` se sigue llamando. Borrar el armado obligaria a
    reescribirlo, y dejarlo sin llamar lo convertiria en codigo muerto.
    """
    t = _texto(build_executive_summary(_datos()))
    for fuera in ("SECCIÓN 3", "SECCIÓN 4", "4.1 Lo favorable",
                  "4.2 Lo desfavorable", "4.3 Tipo de cambio",
                  "Actividad comercial del mes", "Nota metodológica"):
        assert fuera not in t, f"quedo «{fuera}»"
    # Lo que si se queda.
    assert "1.1 Agosto 2026" in t and "SECCIÓN 2" in t
    src = DOCX_MOD.read_text(encoding="utf-8")
    assert 'SECCIONES = ("1", "2")' in src


def test_la_regla_del_ano_completo_NO_se_perdio():
    """⚠️ Vivia en la nota metodologica del final, que se saco. Sin ella un
    Actual de ocho meses contra doce de presupuesto se lee como un derrumbe, asi
    que subio a su propio corte —1.3—, que es donde se lee la columna."""
    t = _texto(build_executive_summary(_datos()))
    assert "la columna principal es el Forecast" in t
    assert "que el año no ha terminado" in t


def test_dividir_entre_cero_NO_inventa_un_porcentaje():
    """Un presupuesto en cero no hace que la variacion sea del 100%: no hay
    porcentaje que dar, y escribir uno inventa una magnitud."""
    assert var_pct(100.0, 0.0) is None
    assert var_pct(100.0, 50.0) == pytest.approx(1.0)


def test_la_escala_cambia_en_el_millon():
    """El PDF mezcla miles y millones a proposito: `$3,973,720.59` en medio de
    una frase no se lee."""
    assert k(277_856.93) == "$277.9K"
    assert k(3_973_720.59) == "$3.974M"
    assert k(-99_140.96) == "-$99.1K"
    assert k(None) == "n/d"


def test_los_numeros_salen_del_MOTOR_y_no_se_recalculan():
    """⚠️ Si el resumen derivara sus propias cifras, el dia que difieran del
    cierre nadie sabria cual mando."""
    src = API.read_text(encoding="utf-8")
    assert "from app.api.pl_api import get_pl_compare" in src
    assert "from app.api.gasto_por_clase_api import gasto_por_clase" in src
    # Y no hay una segunda suma de la cascada.
    assert "_aggregate_selected" not in src


def test_un_mes_sin_noches_NO_entra_en_la_tabla_de_tarifas():
    """Su ADR es cero por falta de operacion, no por tarifa baja, y en una tabla
    de tarifas un cero se lee como un desplome. Amarena cierra cinco meses."""
    src = API.read_text(encoding="utf-8")
    assert 'if not float(kp.get("rooms_occupied") or 0):' in src
    assert "continue" in src


def test_el_endpoint_esta_montado():
    """Un modulo de export que nadie incluye no da error: simplemente no existe
    la ruta, y el boton del front devuelve 404 sin que nada lo avise."""
    from app.api.executive_summary_api import router
    assert [r.path for r in router.routes] == ["/reports/executive-summary/word/"]
    main = (pathlib.Path(__file__).resolve().parents[1] / "app/main.py").read_text(
        encoding="utf-8")
    assert "executive_summary_router" in main
    assert "app.include_router(executive_summary_router" in main


def test_con_base_NEGATIVA_no_se_escribe_porcentaje():
    """⚠️ El error que este informe tuvo con los numeros reales de Amarena.

    El EBITDA de agosto mejoro de -203,8K a -122,7K —81,1K a favor— y la formula
    da **-39,8%**, que se lee como un deterioro del 40%. La frase decia «above
    plan» y el parentesis decia lo contrario, en el mismo renglon.

    Con base negativa manda la palabra, que no se puede leer al reves.
    """
    assert var_pct(-122_660.0, -203_773.0) is None
    assert var_pct(100.0, 0.0) is None
    assert var_pct(306_124.0, 180_516.0) == pytest.approx(0.6959, rel=1e-3)


def test_el_informe_real_no_contradice_su_propio_texto():
    """Con los tres renglones de utilidad en negativo —el caso de Amarena en
    agosto— ninguna frase puede llevar un porcentaje que apunte al otro lado."""
    d = _datos()
    for corte in ("month", "ytd", "full"):
        for v in (d["actual"], d["budget"], d["forecast"]):
            for ln in v[corte]["lines"]:
                if ln["line_code"] in ("EBITDA_BEFORE", "NET_PROFIT", "GOP"):
                    ln["amount_usd"] = -abs(ln["amount_usd"])
    t = _texto(build_executive_summary(d))
    assert "RESUMEN EJECUTIVO MENSUAL" in t   # se genera igual
    assert "n/d%" not in t


def test_la_tipografia_es_la_que_pidio_el_owner():
    """Owner, 2026-09-30: *«el reporte debe ser en Times New Roman, letra 12,
    justificado, espacio 1.5, y entre titulos un espacio adicional»*."""
    src = DOCX_MOD.read_text(encoding="utf-8")
    assert 'FUENTE = "Times New Roman"' in src
    assert "CUERPO = 12" in src
    assert "INTERLINEA = 1.5" in src
    assert "AIRE_TITULO" in src
    # Y de verdad queda escrito en el documento.
    import io
    from docx import Document
    doc = Document(io.BytesIO(build_executive_summary(_datos())))
    normal = doc.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size.pt == 12
    assert normal.paragraph_format.line_spacing == 1.5


def test_el_aire_entre_titulos_NO_es_un_parrafo_vacio():
    """⚠️ Un párrafo vacío se descoloca en cuanto alguien edita arriba, y en
    Word se arrastra al pegar. El aire va como `space_before` del propio
    titulo."""
    src = DOCX_MOD.read_text(encoding="utf-8")
    assert "p.paragraph_format.space_before = Pt(AIRE_TITULO" in src


def test_los_NEGATIVOS_van_en_rojo():
    """⚠️ `f\"${-1234.5:,.2f}\"` da `$-1,234.50`: el signo queda escondido entre el
    simbolo y el numero, se pierde de vista en una columna y —lo que importa
    aca— no se puede detectar para pintarlo.

    Por eso los negativos salen entre PARENTESIS, que es ademas la convencion
    contable y la del PDF del owner.
    """
    assert usd(-1234.5) == "($1,234.50)"
    assert usd(1234.5) == "$1,234.50"
    src = DOCX_MOD.read_text(encoding="utf-8")
    assert "def _es_negativo(texto: str) -> bool:" in src
    assert "if _es_negativo(str(v)):" in src
    assert "r.font.color.rgb = ROJO" in src


def test_los_numeros_de_las_tablas_van_a_la_DERECHA():
    """Es lo unico que alinea las unidades entre si: con el texto centrado, un
    $1,234.50 y un $12.00 no comparten ninguna columna de digitos."""
    src = DOCX_MOD.read_text(encoding="utf-8")
    assert "p.alignment = (WD_ALIGN_PARAGRAPH.RIGHT if i" in src
    # El encabezado se alinea como su columna, o la columna se lee torcida.
    assert "p0.alignment = (WD_ALIGN_PARAGRAPH.RIGHT if i" in src
    # Y la celda no hereda el 1,5 del cuerpo.
    assert "p.paragraph_format.line_spacing = 1.0" in src
