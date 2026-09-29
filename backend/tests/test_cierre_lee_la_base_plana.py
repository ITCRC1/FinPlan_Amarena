# -*- coding: utf-8 -*-
"""El cierre tambien lee la base plana: un archivo, todos los meses.

Owner, 2026-09-28: *«se podra configurar para que en vez de leer el pdf, ahora
lea el excel de datos, en la misma estructura»*.

El PDF trae un mes y hay que subirlo doce veces al ano. La propiedad ya
mantiene un Excel con la misma informacion de todos los meses en una tabla
plana. Este lector devuelve **la misma `LecturaSkill4`** que el de PDF, una
por mes, asi que el calce, la pantalla y el guardado no cambian.
"""
import io
import pathlib

import pytest
from openpyxl import Workbook

from app.importers.datos_planos_room_stats import (
    cuadre_interno, leer_datos_planos,
)
from app.importers.skill4_room_stats_pdf import LecturaSkill4

FRONT = pathlib.Path(__file__).resolve().parents[2] / "frontend"
PAGINA = FRONT / "app/month-end/room-stats/page.tsx"

CAB = ["Mes", "Mes #", "Año", "Tipo de Habitación", "Agencia", "Ing.Hospedaje",
       "Ing.AyB", "Ing.Otros", "Hab.Entradas", "Hab.Estancias", "Cli.Entradas",
       "Cli.Estancias", "Tarifa Prom."]


def libro(filas, cab=CAB, antes=0, hoja="Datos (base plana)"):
    """Un xlsx en memoria con `antes` filas de titulo delante del encabezado."""
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    for _ in range(antes):
        ws.append(["Un titulo cualquiera"])
    ws.append(cab)
    for f in filas:
        ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


FILA_MAR = ["Marzo", 3, 2026, "BEACH FRONT DLXE VILLA", "DIRECTOS",
            407.09, 0, 0, 4, 6, 9, 13, 67.85]
FILA_ABR = ["Abril", 4, 2026, "GARDEN VIEW DLXE VILLA", "BOOKING",
            735.50, 0, 0, 2, 2, 3, 3, 367.75]


# ───────────────────────── Lo que tiene que leer bien ──────────────────────

def test_devuelve_una_lectura_por_mes():
    """Un archivo, todos los meses. Es el punto entero del cambio."""
    lec = leer_datos_planos(libro([FILA_MAR, FILA_ABR]), "x.xlsx", 2026)
    assert [l.month for l in lec] == [3, 4]
    assert all(isinstance(l, LecturaSkill4) for l in lec), \
        "la forma tiene que ser la MISMA que la del PDF o la pantalla se parte"
    assert lec[0].filas[0].hab_estancias == 6
    assert lec[0].filas[0].cli_estancias == 13


def test_entradas_y_estancias_no_se_confunden():
    """⚠️ El error mas caro de este formato, igual que en el PDF.

    Entradas = cuantas reservas llegaron. Estancias = noches. `nights_occupied`
    son las ESTANCIAS de habitaciones y `pax` las ESTANCIAS de clientes.
    Tomar las entradas da un ADR y un ratio que parecen razonables y estan
    mal: 407.09/4 = 101.77 en vez de 407.09/6 = 67.85, que es justo lo que el
    propio archivo imprime en «Tarifa Prom.».
    """
    lec = leer_datos_planos(libro([FILA_MAR]), "x.xlsx", 2026)[0]
    f = lec.filas[0]
    assert f.hab_entradas == 4 and f.hab_estancias == 6
    assert f.cli_entradas == 9 and f.cli_estancias == 13
    t = lec.por_tipo()[0]
    assert t["nights_occupied"] == 6, "nights_occupied tomo las entradas"
    assert t["pax"] == 13, "pax tomo las entradas de clientes"
    assert t["adr"] == pytest.approx(67.85, abs=0.01)


def test_las_columnas_se_buscan_por_nombre_no_por_posicion():
    """Insertar una columna al medio del Excel no puede mover un numero."""
    cab = CAB[:5] + ["Comentario"] + CAB[5:]
    filas = [FILA_MAR[:5] + ["lo que sea"] + FILA_MAR[5:]]
    lec = leer_datos_planos(libro(filas, cab=cab), "x.xlsx", 2026)[0]
    assert lec.filas[0].ingreso_hospedaje == pytest.approx(407.09)
    assert lec.filas[0].hab_estancias == 6


def test_el_encabezado_puede_no_estar_en_la_primera_fila():
    """Los archivos del grupo traen titulo y subtitulo arriba de la tabla."""
    lec = leer_datos_planos(libro([FILA_MAR], antes=4), "x.xlsx", 2026)
    assert lec[0].filas[0].agencia == "DIRECTOS"


def test_se_elige_la_hoja_que_tiene_la_tabla():
    """⚠️ El libro real trae diez hojas y la base plana es la ultima.

    Pedirle a la persona que la renombre o que la ponga primera seria un paso
    mas que se puede equivocar, y la hoja correcta se reconoce sola por sus
    columnas.
    """
    wb = Workbook()
    wb.active.title = "Resumen"
    wb.active.append(["Un resumen cualquiera", 1, 2, 3])
    ws = wb.create_sheet("Datos (base plana)")
    ws.append(CAB)
    ws.append(FILA_MAR)
    buf = io.BytesIO()
    wb.save(buf)
    lec = leer_datos_planos(buf.getvalue(), "x.xlsx", 2026)
    assert len(lec) == 1 and lec[0].filas[0].agencia == "DIRECTOS"


def test_el_mes_sale_del_numero_y_si_no_del_nombre():
    """El numero no depende del idioma; el nombre es el respaldo."""
    sin_num = ["Julio", None, 2026, "X", "Y", 100, 0, 0, 1, 2, 1, 2, 50]
    lec = leer_datos_planos(libro([sin_num]), "x.xlsx", 2026)
    assert lec[0].month == 7


# ─────────────────────── Lo que tiene que rechazar ─────────────────────────

def test_rechaza_un_archivo_sin_encabezado():
    wb = Workbook()
    wb.active.append(["esto", "no", "es", "la", "tabla"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="encabezado"):
        leer_datos_planos(buf.getvalue(), "x.xlsx", 2026)


def test_rechaza_un_ano_que_no_es_el_del_escenario():
    """⚠️ Mismo criterio que el PDF. Sin esto, la base de 2025 entraria como
    2026 —mismo mes, otro ano— y el reporte no tendria como notarlo."""
    fila = list(FILA_MAR)
    fila[2] = 2025
    with pytest.raises(ValueError, match="2025"):
        leer_datos_planos(libro([fila]), "x.xlsx", 2026)


def test_una_fila_de_total_dentro_de_la_base_se_descarta_y_se_avisa():
    """⚠️ El modo de falla que duplica todo sin que nada avise.

    Si alguien copia la tabla del reporte con sus filas de TOTAL, esas filas
    suman de nuevo lo que el detalle ya trae: el mes queda al doble y sigue
    cuadrando consigo mismo.
    """
    total = ["Marzo", 3, 2026, "TOTAL BEACH FRONT DLXE VILLA", "TOTAL",
             2481.43, 0, 0, 15, 21, 30, 42, 118.16]
    lec = leer_datos_planos(libro([FILA_MAR, total]), "x.xlsx", 2026)
    assert len(lec[0].filas) == 1, "la fila de total entro como si fuera detalle"
    avisos = getattr(lec[0], "avisos_del_archivo", [])
    assert any("total" in a.lower() for a in avisos), "se descarto en silencio"


def test_una_fila_repetida_se_suma_pero_se_avisa():
    """Puede ser legitimo (dos bloques del PMS) o un copiar-pegar duplicado.
    El total cuadra en los dos casos, asi que hay que decirlo."""
    lec = leer_datos_planos(libro([FILA_MAR, FILA_MAR]), "x.xlsx", 2026)
    assert lec[0].por_tipo()[0]["nights_occupied"] == 12
    avisos = getattr(lec[0], "avisos_del_archivo", [])
    assert any("ya aparecio" in a for a in avisos)


def test_rechaza_un_archivo_sin_filas_de_datos():
    with pytest.raises(ValueError, match="ninguna fila"):
        leer_datos_planos(libro([]), "x.xlsx", 2026)


# ──────────────── El cuadre que este camino SI puede hacer ─────────────────

def test_el_cuadre_interno_detecta_una_celda_editada_a_mano():
    """⚠️ Este camino NO tiene cuadre contra una segunda fuente.

    El PDF cierra con un resumen del hotel y se verifica contra el por tres
    vias. La base plana es solo el detalle y se cree entera. Lo unico
    verificable adentro es la tarifa impresa: si alguien toca el ingreso o
    las noches, deja de ser ingreso/noches.
    """
    mal = list(FILA_MAR)
    mal[5] = 500.00            # se edito el ingreso y no la tarifa
    lec = leer_datos_planos(libro([mal]), "x.xlsx", 2026)[0]
    avisos = cuadre_interno(lec)
    assert avisos and "tarifa" in avisos[0]


def test_el_cuadre_interno_detecta_ingreso_sin_noches():
    fila = ["Marzo", 3, 2026, "X", "Y", 900.0, 0, 0, 0, 0, 0, 0, 0]
    lec = leer_datos_planos(libro([fila]), "x.xlsx", 2026)[0]
    assert any("cero noches" in a for a in cuadre_interno(lec))


def test_el_cuadre_interno_detecta_mas_entradas_que_noches():
    """Una entrada genera al menos una noche: al reves es imposible."""
    fila = ["Marzo", 3, 2026, "X", "Y", 100.0, 0, 0, 9, 2, 9, 2, 50.0]
    lec = leer_datos_planos(libro([fila]), "x.xlsx", 2026)[0]
    assert any("entradas" in a for a in cuadre_interno(lec))


def test_el_resumen_del_hotel_queda_en_cero_y_no_se_inventa():
    """⚠️ La base plana no trae capacidad, disponibles ni bloqueadas.

    Rellenarlas con `capacidad x dias` fabricaria justamente la cifra contra
    la que habria que contrastar, y la pantalla las dibujaria como si el
    archivo las hubiera dicho.
    """
    lec = leer_datos_planos(libro([FILA_MAR]), "x.xlsx", 2026)[0]
    r = lec.resumen
    assert r.habitaciones_disponibles == 0
    assert r.habitaciones_bloqueadas == 0
    assert r.ingreso_total_hotel == 0
    assert r.dias == 31, "los dias del mes si se saben, salen del calendario"


# ───────────────────────── El endpoint y la pantalla ───────────────────────

def test_el_endpoint_existe_y_no_guarda_nada():
    """Leer no es importar. El archivo se descarta al terminar la peticion."""
    import inspect

    from app.api import room_stats_pdf_api as api
    assert hasattr(api, "leer_excel_room_stats")
    rutas = [r for r in api.router.routes
             if getattr(r, "path", "").endswith("/room-stats/leer-excel/")]
    assert rutas and "POST" in rutas[0].methods
    src = inspect.getsource(api.leer_excel_room_stats)
    for escribe in ("db.add", "db.commit", "db.merge", "delete("):
        assert escribe not in src, f"el lector {escribe} — tiene que leer y nada mas"


def test_los_dos_lectores_arman_el_mes_con_la_misma_funcion():
    """⚠️ El calce de categorias y el de canales no pueden divergir segun de
    que archivo vino el mes. Si divergieran, el mismo dato entraria distinto
    segun el camino y nadie lo notaria hasta comparar dos cargas."""
    import inspect

    from app.api import room_stats_pdf_api as api
    assert "_armar_mes(" in inspect.getsource(api.leer_pdf_room_stats)
    assert "_armar_mes(" in inspect.getsource(api.leer_excel_room_stats)


def test_la_pantalla_acepta_los_dos_formatos():
    pag = PAGINA.read_text(encoding="utf-8")
    assert "leerExcelRoomStats" in pag
    assert ".xlsx" in pag and ".csv" in pag
    assert "/\\.(xlsx|xlsm|xls|csv)$/i.test(f.name)" in pag, \
        "no se decide el lector por la extension"


def test_cambiar_de_mes_no_obliga_a_volver_a_subir():
    """⚠️ Sin esto, un Excel de seis meses habia que subirlo seis veces — el
    mismo trabajo que este camino viene a sacar."""
    pag = PAGINA.read_text(encoding="utf-8")
    assert "const [delArchivo" in pag
    assert "const otro = delArchivo[mesSel];" in pag


def test_guardar_todos_usa_el_mismo_guardado_que_uno():
    """Si el guardado masivo armara las filas por su cuenta, un mes entraria
    distinto segun que boton se apreto."""
    pag = PAGINA.read_text(encoding="utf-8")
    assert "async function guardarUno" in pag
    cuerpo = pag[pag.index("async function guardarTodos"):pag.index("async function guardar()")]
    assert "await guardarUno(m)" in cuerpo
    assert "for (const m of meses)" in cuerpo, \
        "los meses se mandan en paralelo: cada guardado reemplaza el mes entero"


def test_el_calce_se_siembra_con_todos_los_meses():
    """⚠️ Agosto trae la categoria Accesible y marzo no. Sembrando el calce
    solo con el mes visible, «Guardar los 6» mandaba esa fila sin categoria."""
    pag = PAGINA.read_text(encoding="utf-8")
    assert "r.meses.flatMap(m =>" in pag
    assert "const faltanCalce = useMemo" in pag
